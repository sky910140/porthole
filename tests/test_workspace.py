from __future__ import annotations

import os
from pathlib import Path

import pytest

from project_mcp.workspace import Workspace


def test_workspace_info_and_file_listing_are_bounded_and_filter_secrets(tmp_path: Path) -> None:
    (tmp_path / "src").mkdir()
    (tmp_path / "src" / "app.py").write_text("print('ok')\n", encoding="utf-8")
    (tmp_path / ".env").write_text("TOKEN=secret", encoding="utf-8")
    (tmp_path / "node_modules").mkdir()
    (tmp_path / "node_modules" / "leak.js").write_text("secret", encoding="utf-8")

    workspace = Workspace(tmp_path, "demo")

    info = workspace.workspace_info()
    listing = workspace.list_files(limit=1)
    assert info["project_id"] == "demo"
    assert info["root"] == "."
    assert listing == {
        "project_id": "demo",
        "path": ".",
        "files": [{"path": "src/app.py", "size": 13, "read_as": "text_candidate"}],
        "offset": 0,
        "limit": 1,
        "source": "disk",
        "has_more": False,
        "next_offset": None,
        "truncation_reason": None,
        "truncated": False,
    }


@pytest.mark.parametrize(
    "path",
    ["../outside.txt", "/absolute.txt", r"C:\\Windows\\win.ini", "file.txt:stream", "CON"],
)
def test_resolve_file_rejects_traversal_absolute_ads_and_device_paths(
    tmp_path: Path, path: str
) -> None:
    workspace = Workspace(tmp_path, "demo")
    with pytest.raises((ValueError, PermissionError)):
        workspace.resolve_file(path)


def test_resolve_file_rejects_symlink_traversal(tmp_path: Path) -> None:
    outside = tmp_path.parent / f"{tmp_path.name}-outside"
    outside.mkdir()
    (outside / "secret.txt").write_text("hidden", encoding="utf-8")
    link = tmp_path / "linked"
    try:
        link.symlink_to(outside, target_is_directory=True)
    except OSError as exc:
        pytest.skip(f"symlink creation unavailable: {exc}")

    with pytest.raises(PermissionError):
        Workspace(tmp_path, "demo").resolve_file("linked/secret.txt")


def test_read_file_returns_fresh_text_metadata_and_rejects_binary_or_large_files(
    tmp_path: Path,
) -> None:
    source = tmp_path / "hello.py"
    source.write_text("one\ntwo\nthree\n", encoding="utf-8")
    workspace = Workspace(tmp_path, "demo")

    first = workspace.read_file("hello.py", 2, 3)
    source.write_text("changed\ntwo\nthree\n", encoding="utf-8")
    os.utime(source, None)
    second = workspace.read_file("hello.py", 1, 1)

    assert first["lines"] == [{"line": 2, "text": "two"}, {"line": 3, "text": "three"}]
    assert first["sha256"] != second["sha256"]
    assert second["lines"] == [{"line": 1, "text": "changed"}]
    assert isinstance(second["mtime_ns"], int)

    (tmp_path / "binary.bin").write_bytes(b"hello\x00world")
    (tmp_path / "large.txt").write_bytes(b"x" * (1024 * 1024 + 1))
    with pytest.raises(ValueError, match="binary"):
        workspace.read_file("binary.bin")
    with pytest.raises(ValueError, match="too large"):
        workspace.read_file("large.txt")


def test_search_code_is_fixed_string_filtered_and_limited(tmp_path: Path) -> None:
    (tmp_path / "a.py").write_text("a+b\na+b\n", encoding="utf-8")
    (tmp_path / "b.py").write_text("a.*b\n", encoding="utf-8")
    (tmp_path / "credentials.json").write_text('a+b: "secret"', encoding="utf-8")

    result = Workspace(tmp_path, "demo").search_code("a+b", limit=1)

    assert result["matches"] == [{"path": "a.py", "line": 1, "text": "a+b"}]
    assert result["truncated"] is True
    with pytest.raises(ValueError):
        Workspace(tmp_path, "demo").search_code("")


@pytest.mark.parametrize(
    "path",
    [
        ".local/token.txt",
        ".ssh/id_ed25519",
        ".codex/auth.json",
        ".agents/state.json",
        "config/local.json",
        "tokens.json",
        "server.pem",
        "client.key",
        "certificate.p12",
        "certificate.pfx",
        "nested/my_private_key.txt",
        "nested/config/local.json",
    ],
)
def test_sensitive_paths_are_excluded_from_read_and_search(tmp_path: Path, path: str) -> None:
    target = tmp_path / Path(path)
    target.parent.mkdir(parents=True, exist_ok=True)
    target.write_text("unique-secret-marker", encoding="utf-8")
    workspace = Workspace(tmp_path, "demo")

    with pytest.raises(PermissionError):
        workspace.read_file(path)
    assert workspace.search_code("unique-secret-marker")["matches"] == []


@pytest.mark.parametrize("path", [".env.", ".env ", "folder./file.txt", "folder /file.txt"])
def test_windows_trailing_dot_or_space_aliases_are_rejected(tmp_path: Path, path: str) -> None:
    with pytest.raises(ValueError):
        Workspace(tmp_path, "demo").resolve_file(path)


def test_list_files_reports_scan_truncation(monkeypatch: pytest.MonkeyPatch, tmp_path: Path) -> None:
    for index in range(3):
        (tmp_path / f"{index}.txt").write_text(str(index), encoding="utf-8")
    monkeypatch.setattr("project_mcp.workspace.MAX_SCAN_FILES", 2)

    result = Workspace(tmp_path, "demo").list_files(limit=10)

    assert len(result["files"]) == 2
    assert result["truncated"] is True
    assert result["has_more"] is False
    assert result["next_offset"] is None
    assert result["truncation_reason"] == "scan_limit"


def test_list_files_distinguishes_next_page_from_scan_limit(tmp_path: Path) -> None:
    for index in range(3):
        (tmp_path / f"{index}.txt").write_text(str(index), encoding="utf8")
    workspace = Workspace(tmp_path, "demo")
    first = workspace.list_files(limit=2)
    second = workspace.list_files(limit=2, offset=first["next_offset"])
    assert first["has_more"] is True
    assert first["truncation_reason"] == "page_limit"
    assert first["next_offset"] == 2
    assert [item["path"] for item in second["files"]] == ["2.txt"]
    assert second["has_more"] is False


def test_read_files_is_independent_per_item_and_has_total_budget(monkeypatch, tmp_path: Path) -> None:
    (tmp_path / "a.txt").write_text("alpha\n", encoding="utf8")
    (tmp_path / "b.txt").write_text("beta\n", encoding="utf8")
    (tmp_path / ".env").write_text("secret", encoding="utf8")
    workspace = Workspace(tmp_path, "demo")
    result = workspace.read_files([
        {"path": "a.txt", "start_line": 1, "end_line": 10},
        {"path": ".env"},
        {"path": "missing.txt"},
    ])
    assert result["results"][0]["ok"] is True
    assert result["results"][1]["error"]["code"] == "PATH_FORBIDDEN"
    assert result["results"][2]["error"]["code"] == "PATH_FORBIDDEN"

    monkeypatch.setattr("project_mcp.workspace.MAX_BATCH_READ_BYTES", 5)
    limited = workspace.read_files([{"path": "a.txt"}, {"path": "b.txt"}])
    assert limited["results"][1]["error"]["code"] == "QUOTA_EXCEEDED"


def test_scope_preview_reports_exclusions_without_caching_authorization(tmp_path: Path) -> None:
    from project_mcp.policy import ProjectPolicy

    (tmp_path / "src").mkdir()
    (tmp_path / "src" / "app.py").write_text("ok", encoding="utf8")
    (tmp_path / ".env").write_text("secret", encoding="utf8")
    policy = ProjectPolicy(exclude_paths=["src/private/**"])
    workspace = Workspace(tmp_path, "demo", policy=policy)
    preview = workspace.preview_scope()
    assert preview["accessible_files"] == 1
    assert preview["excluded_by_reason"]["sensitive_path"] >= 1
    assert preview["scan_complete"] is True


def test_scope_preview_lists_allowed_files_without_contents_or_excluded_paths(tmp_path):
    from project_mcp.policy import ProjectPolicy

    (tmp_path / "src").mkdir()
    (tmp_path / "src" / "说明.txt").write_text("hello", encoding="utf-8")
    (tmp_path / "estimate.csv").write_text("a,b", encoding="utf-8")
    (tmp_path / "drawing.pdf").write_bytes(b"%PDF")
    (tmp_path / ".env").write_text("do-not-share", encoding="utf-8")
    (tmp_path / "private.txt").write_text("do-not-share", encoding="utf-8")
    (tmp_path / "large.txt").write_bytes(b"x" * (1024 * 1024 + 1))
    workspace = Workspace(tmp_path, "demo", policy=ProjectPolicy(exclude_paths=["private.txt"]))
    preview = workspace.preview_scope()
    assert preview["files"] == [
        {"path": "drawing.pdf", "size": 4, "read_as": "unsupported"},
        {"path": "estimate.csv", "size": 3, "read_as": "table"},
        {"path": "src/说明.txt", "size": 5, "read_as": "text_candidate"},
    ]
    assert preview["accessible_files"] == 3
    assert preview["excluded_by_reason"] == {
        "sensitive_path": 1, "configured_exclusion": 1, "file_too_large": 1,
    }
    assert preview["files_truncated"] is False
    assert "do-not-share" not in str(preview)
    assert str(tmp_path) not in str(preview)
    (tmp_path / "estimate.csv").unlink()
    assert "estimate.csv" not in str(workspace.preview_scope())


def test_scope_preview_bounds_file_details_without_truncating_scan(tmp_path):
    for index in range(205):
        (tmp_path / f"{index:03}.txt").write_text("ok", encoding="utf-8")
    result = Workspace(tmp_path, "demo").preview_scope()
    assert result["accessible_files"] == 205
    assert len(result["files"]) == 200
    assert result["files_truncated"] is True
    assert result["scan_complete"] is True
    assert result["truncation_reason"] is None


def test_scope_preview_empty_folder_and_scan_limit_are_explicit(tmp_path, monkeypatch):
    workspace = Workspace(tmp_path, "demo")
    result = workspace.preview_scope()
    assert result["files"] == []
    assert result["files_truncated"] is False
    for index in range(3):
        (tmp_path / f"{index}.txt").write_text("ok", encoding="utf-8")
    monkeypatch.setattr("project_mcp.workspace.MAX_SCAN_FILES", 2)
    result = workspace.preview_scope()
    assert len(result["files"]) == 2
    assert result["scan_complete"] is False
    assert result["truncation_reason"] == "scan_limit"


def test_scope_preview_unavailable_directory_does_not_claim_complete(tmp_path, monkeypatch):
    original = os.scandir
    (tmp_path / "locked").mkdir()

    def scan(directory):
        if Path(directory).name == "locked":
            raise PermissionError("unavailable")
        return original(directory)

    monkeypatch.setattr(os, "scandir", scan)
    result = Workspace(tmp_path, "demo").preview_scope()
    assert result["scan_complete"] is False
    assert result["truncation_reason"] == "unavailable"
    assert result["excluded_by_reason"] == {"unavailable": 1}


def test_scope_preview_does_not_disclose_files_when_project_is_paused(tmp_path):
    from project_mcp.policy import ProjectPolicy

    (tmp_path / "private-name.txt").write_text("hello", encoding="utf-8")
    workspace = Workspace(tmp_path, "demo", policy=ProjectPolicy(paused=True))
    with pytest.raises(PermissionError, match="paused"):
        workspace.preview_scope()


def test_custom_excluded_paths_are_hidden_from_every_workspace_tool(tmp_path: Path) -> None:
    state_dir = tmp_path / "runtime"
    state_dir.mkdir()
    secret = state_dir / "session.json"
    secret.write_text("custom-runtime-secret", encoding="utf-8")
    config = tmp_path / "assistant.json"
    config.write_text("custom-config-secret", encoding="utf-8")
    workspace = Workspace(tmp_path, "demo", excluded_paths=[state_dir, config])

    assert workspace.list_files()["files"] == []
    assert workspace.search_code("custom-runtime-secret")["matches"] == []
    assert workspace.search_code("custom-config-secret")["matches"] == []
    with pytest.raises(PermissionError):
        workspace.read_file("runtime/session.json")
    with pytest.raises(PermissionError):
        workspace.read_file("assistant.json")


def test_read_table_csv_is_bounded_and_reports_more_rows_and_columns(tmp_path: Path) -> None:
    (tmp_path / "estimate.csv").write_text(
        "项目,数量,单价\n电缆,2,10\n支架,3,20\n", encoding="utf-8-sig"
    )
    result = Workspace(tmp_path, "demo").read_table(
        "estimate.csv", start_row=2, limit=1, max_columns=2
    )
    assert result["rows"] == [{"row": 2, "cells": ["电缆", "2"]}]
    assert result["has_more"] is True
    assert result["next_row"] == 3
    assert result["columns_truncated"] is True


def test_read_table_xlsx_selects_sheet_and_never_evaluates_formulas(tmp_path: Path) -> None:
    from openpyxl import Workbook

    workbook = Workbook()
    workbook.active.title = "说明"
    detail = workbook.create_sheet("清单")
    detail.append(["名称", "金额"])
    detail.append(["电缆", "=1+2"])
    workbook.save(tmp_path / "estimate.xlsx")

    result = Workspace(tmp_path, "demo").read_table("estimate.xlsx", sheet="清单")
    assert result["sheet"] == "清单"
    assert result["sheets"] == ["说明", "清单"]
    assert result["rows"][0]["cells"] == ["名称", "金额"]
    assert result["rows"][1]["cells"] == ["电缆", None]
    assert Workspace(tmp_path, "demo").read_table(
        "estimate.xlsx", sheet="清单", start_row=100
    )["rows"] == []


def test_read_table_marks_nonempty_xlsx_columns_beyond_preview(tmp_path: Path) -> None:
    from openpyxl import Workbook

    workbook = Workbook()
    workbook.active.cell(row=1, column=22, value="late value")
    workbook.save(tmp_path / "wide.xlsx")
    result = Workspace(tmp_path, "demo").read_table("wide.xlsx", max_columns=20)
    assert result["columns_truncated"] is True
    assert result["truncation_reason"] == "column_limit"


def test_read_table_rejects_secret_unsupported_and_oversized_files(tmp_path: Path) -> None:
    (tmp_path / ".env").write_text("secret", encoding="utf-8")
    (tmp_path / "drawing.pdf").write_bytes(b"%PDF")
    (tmp_path / "large.csv").write_bytes(b"a" * (1024 * 1024 + 1))
    workspace = Workspace(tmp_path, "demo")
    with pytest.raises(PermissionError):
        workspace.read_table(".env")
    with pytest.raises(ValueError, match="CSV or XLSX"):
        workspace.read_table("drawing.pdf")
    with pytest.raises(ValueError, match="too large"):
        workspace.read_table("large.csv")


def test_read_table_rejects_highly_expanded_xlsx_and_limits_response(monkeypatch, tmp_path: Path) -> None:
    from zipfile import ZIP_DEFLATED, ZipFile

    with ZipFile(tmp_path / "bomb.xlsx", "w", compression=ZIP_DEFLATED) as archive:
        archive.writestr("xl/blob.xml", "A" * (16 * 1024 * 1024 + 1))
    workspace = Workspace(tmp_path, "demo")
    with pytest.raises(ValueError, match="expanded"):
        workspace.read_table("bomb.xlsx")

    (tmp_path / "small.csv").write_text("abcdef\nghijkl\n", encoding="utf-8")
    monkeypatch.setattr("project_mcp.table_read.MAX_TABLE_OUTPUT_BYTES", 5)
    with pytest.raises(ValueError, match="response metadata exceeds limit"):
        workspace.read_table("small.csv")


def test_read_table_reports_malformed_sheet_as_a_safe_input_error(tmp_path: Path) -> None:
    from zipfile import ZipFile

    from openpyxl import Workbook

    source = tmp_path / "source.xlsx"
    workbook = Workbook()
    workbook.active.append(["value"])
    workbook.save(source)
    with ZipFile(source) as original, ZipFile(tmp_path / "broken.xlsx", "w") as broken:
        for member in original.infolist():
            payload = original.read(member.filename)
            if member.filename == "xl/worksheets/sheet1.xml":
                payload = b"<worksheet><sheetData><row>"
            broken.writestr(member, payload)
    with pytest.raises(ValueError, match="invalid|parser"):
        Workspace(tmp_path, "demo").read_table("broken.xlsx")


def test_read_table_caps_complete_json_response_and_csv_growth(monkeypatch, tmp_path: Path) -> None:
    import json
    from types import SimpleNamespace

    from project_mcp.table_read import MAX_TABLE_OUTPUT_BYTES, read_table

    data = tmp_path / "many.csv"
    data.write_text((",".join(["x" * 15] * 50) + "\n") * 100, encoding="utf-8")
    result = Workspace(tmp_path, "demo").read_table("many.csv", limit=100, max_columns=50)
    assert len(json.dumps(result, ensure_ascii=False).encode("utf-8")) <= MAX_TABLE_OUTPUT_BYTES
    assert result["has_more"] is True

    data.write_bytes(b"a" * (1024 * 1024 + 100))
    original_stat = Path.stat

    def stale_stat(file_path, *args, **kwargs):
        value = original_stat(file_path, *args, **kwargs)
        return SimpleNamespace(st_size=1) if file_path == data else value

    monkeypatch.setattr(Path, "stat", stale_stat)
    with pytest.raises(ValueError, match="too large"):
        read_table(data, project_id="demo", relative_path="many.csv")


def test_read_table_can_return_a_full_width_unicode_row(tmp_path: Path) -> None:
    (tmp_path / "unicode.csv").write_text(
        ",".join(["😀" * 500] * 50) + "\n", encoding="utf-8"
    )
    result = Workspace(tmp_path, "demo").read_table(
        "unicode.csv", limit=1, max_columns=50
    )
    assert len(result["rows"]) == 1
    assert len(result["rows"][0]["cells"]) == 50
    assert result["truncation_reason"] == "cell_length"


def test_file_listing_discloses_table_and_unsupported_binary_reading(tmp_path: Path) -> None:
    (tmp_path / "notes.md").write_text("hello", encoding="utf-8")
    (tmp_path / "estimate.csv").write_text("name\n", encoding="utf-8")
    (tmp_path / "estimate.xlsx").write_bytes(b"placeholder")
    (tmp_path / "drawing.pdf").write_bytes(b"%PDF")
    files = {item["path"]: item for item in Workspace(tmp_path, "demo").list_files()["files"]}
    assert files["notes.md"]["read_as"] == "text_candidate"
    assert files["estimate.csv"]["read_as"] == "table"
    assert files["estimate.xlsx"]["read_as"] == "table"
    assert files["drawing.pdf"]["read_as"] == "unsupported"
