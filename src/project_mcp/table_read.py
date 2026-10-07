"""Bounded, read-only CSV and XLSX previews for explicitly authorized files."""

from __future__ import annotations

import csv
import io
import json
from itertools import islice
from pathlib import Path
from xml.etree.ElementTree import ParseError
from zipfile import BadZipFile, ZipFile

import defusedxml  # noqa: F401 - enables openpyxl's guarded XML parser
from defusedxml.common import DefusedXmlException
from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException

MAX_TABLE_FILE_BYTES = 1024 * 1024
MAX_TABLE_EXPANDED_BYTES = 16 * 1024 * 1024
MAX_TABLE_OUTPUT_BYTES = 64 * 1024
MAX_TABLE_CELL_CHARS = 100
MAX_TABLE_MEMBERS = 512


def _cell(value: object) -> tuple[str | None, bool]:
    if value is None:
        return None, False
    rendered = str(value)
    return rendered[:MAX_TABLE_CELL_CHARS], len(rendered) > MAX_TABLE_CELL_CHARS


def _check_archive(snapshot: io.BytesIO) -> None:
    try:
        with ZipFile(snapshot) as archive:
            members = archive.infolist()
            if len(members) > MAX_TABLE_MEMBERS or sum(
                item.file_size for item in members
            ) > MAX_TABLE_EXPANDED_BYTES:
                raise ValueError("XLSX expanded content is too large")
            if any(item.file_size > MAX_TABLE_EXPANDED_BYTES for item in members):
                raise ValueError("XLSX expanded content is too large")
    except BadZipFile:
        raise ValueError("invalid XLSX file") from None


def read_table(
    file_path: Path, *, project_id: str, relative_path: str,
    sheet: str | None = None, start_row: int = 1, limit: int = 50,
    max_columns: int = 20,
) -> dict[str, object]:
    if not isinstance(start_row, int) or not 1 <= start_row <= 1_000_000:
        raise ValueError("start_row must be between 1 and 1000000")
    if not isinstance(limit, int) or not 1 <= limit <= 100:
        raise ValueError("limit must be between 1 and 100")
    if not isinstance(max_columns, int) or not 1 <= max_columns <= 50:
        raise ValueError("max_columns must be between 1 and 50")
    if sheet is not None and (not isinstance(sheet, str) or not sheet or len(sheet) > 100):
        raise ValueError("invalid sheet name")
    if file_path.stat().st_size > MAX_TABLE_FILE_BYTES:
        raise ValueError("file is too large")
    with file_path.open("rb") as source:
        raw = source.read(MAX_TABLE_FILE_BYTES + 1)
    if len(raw) > MAX_TABLE_FILE_BYTES:
        raise ValueError("file is too large")

    suffix = file_path.suffix.lower()
    if suffix == ".csv":
        if sheet is not None:
            raise ValueError("CSV files do not have sheets")
        if b"\x00" in raw:
            raise ValueError("binary CSV file is not supported")
        try:
            text = raw.decode("utf-8-sig")
        except UnicodeDecodeError:
            raise ValueError("CSV must be UTF-8") from None
        reader = csv.reader(io.StringIO(text, newline=""))
        selected = islice(enumerate(reader, 1), start_row - 1, start_row + limit)
        sheet_name = None
        sheets: list[str] = []
        workbook = None
    elif suffix == ".xlsx":
        _check_archive(io.BytesIO(raw))
        try:
            workbook = load_workbook(
                io.BytesIO(raw), read_only=True, data_only=True, keep_links=False
            )
        except (OSError, ValueError, BadZipFile, KeyError, ParseError,
                DefusedXmlException, InvalidFileException):
            raise ValueError("invalid XLSX file") from None
        sheets = workbook.sheetnames[:50]
        sheet_name = sheet or (sheets[0] if sheets else None)
        if sheet_name not in workbook.sheetnames:
            workbook.close()
            raise ValueError("unknown sheet; use one of the returned sheet names")
        worksheet = workbook[sheet_name]
        # A narrow preview must disclose columns beyond its visible range.
        extra_sheet_columns = (worksheet.max_column or 0) > max_columns
        if worksheet.max_row is not None and start_row > worksheet.max_row:
            selected = iter(())
        else:
            last_row = min(start_row + limit, worksheet.max_row or start_row + limit)
            selected = enumerate(
                worksheet.iter_rows(
                    min_row=start_row, max_row=last_row,
                    min_col=1, max_col=min(worksheet.max_column or 1, max_columns + 1),
                    values_only=True,
                ), start_row,
            )
    else:
        raise ValueError("read_table supports CSV or XLSX only")

    rows: list[dict[str, object]] = []
    output_bytes = 0
    columns_truncated = suffix == ".xlsx" and extra_sheet_columns
    cells_truncated = False
    has_more = False
    output_limited = False
    try:
        for number, raw_row in selected:
            if len(rows) >= limit:
                has_more = True
                break
            if len(raw_row) > max_columns and any(
                item not in (None, "") for item in raw_row[max_columns:]
            ):
                columns_truncated = True
            cells: list[str | None] = []
            for value in raw_row[:max_columns]:
                rendered, limited = _cell(value)
                cells.append(rendered)
                cells_truncated |= limited
            row_bytes = len(json.dumps(cells, ensure_ascii=False).encode("utf-8"))
            if output_bytes + row_bytes > MAX_TABLE_OUTPUT_BYTES:
                output_limited = True
                has_more = True
                break
            output_bytes += row_bytes
            rows.append({"row": number, "cells": cells})
    except (csv.Error, ValueError, ParseError, DefusedXmlException, BadZipFile):
        raise ValueError("table content is invalid or exceeds parser limits") from None
    finally:
        if workbook is not None:
            workbook.close()

    reason = (
        "output_limit" if output_limited else "page_limit" if has_more
        else "column_limit" if columns_truncated
        else "cell_length" if cells_truncated else None
    )
    result = {
        "project_id": project_id, "path": relative_path, "source": "disk",
        "format": suffix[1:], "sheet": sheet_name, "sheets": sheets,
        "start_row": start_row, "limit": limit, "max_columns": max_columns,
        "rows": rows, "has_more": has_more,
        "next_row": start_row + len(rows) if has_more and rows else None,
        "columns_truncated": columns_truncated,
        "truncation_reason": reason, "truncated": reason is not None,
    }
    while len(json.dumps(result, ensure_ascii=False).encode("utf-8")) > MAX_TABLE_OUTPUT_BYTES:
        if not rows:
            raise ValueError("table response metadata exceeds limit")
        rows.pop()
        result["has_more"] = True
        result["next_row"] = start_row + len(rows) if rows else None
        result["truncation_reason"] = "output_limit"
        result["truncated"] = True
    return result
